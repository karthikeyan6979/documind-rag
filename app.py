from fastapi import FastAPI, UploadFile, HTTPException
from pydantic import BaseModel
import chromadb
import requests
import uuid
import time
import pdfplumber
import io
from docx import Document
import json
from fastapi.responses import StreamingResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from bs4 import BeautifulSoup
import ipaddress
import socket
from urllib.parse import urlparse
from fastapi.middleware.cors import CORSMiddleware
import re
from collections import Counter


app = FastAPI()

app.add_middleware(CORSMiddleware,allow_origins=["*"],allow_methods=["*"],allow_headers=["*"],) # type: ignore


chat_history = {}
client = chromadb.PersistentClient(path = "./chroma_db")
collection = client.get_or_create_collection(name="documents", metadata = {"hnsw:space": "cosine"})

def chunking(text, chunk_size = 500, overlap_sentence = 1):
    sentences = text.split(".")
    current_chunk = []
    current_length = 0
    chunks = []
    for sentence in sentences:
        sentence = sentence.strip()
        if not sentence:
            continue
        if len(sentence) + current_length > chunk_size and current_chunk:
            chunks.append(".".join(current_chunk) + ".")
            current_chunk = current_chunk[-overlap_sentence:]
            current_length = 0 
            for i in current_chunk:
                current_length += len(i)
        current_chunk.append(sentence)
        current_length += len(sentence)
    if current_chunk:
        chunks.append(".".join(current_chunk) + ".")

    return chunks

def get_embeddings_batch(texts: list[str]):
    response = requests.post(
        "http://localhost:11434/api/embed",
        json={"model": "nomic-embed-text", "input": texts}
    )
    data = response.json()
    embeddings = data.get("embeddings")
    
    if not embeddings or len(embeddings) != len(texts):
        raise ValueError(f"Mismatch or empty response: expected {len(texts)}, got {len(embeddings) if embeddings else 0}")
    
    for i, e in enumerate(embeddings):
        if not e:
            raise ValueError(f"Empty embedding at index {i} for text: {texts[i][:80]!r}")
    
    return embeddings

def get_embedding_single(text: str):

    return get_embeddings_batch([text])[0]

def rewrite_query(question, history):
    if not history:
        return question
    history_text = "\n".join([f"Q: {q}\nA: {a}" for q, a in history])
    prompt = f""" Given the conversation history and a new question, decide if the new question depends on the history (e.g. uses pronouns like "it", "that", "this", or is otherwise incomplete without context).

    If it depends on history, rewrite it to be a standalone question.
    If it does NOT depend on history (it's already a complete, standalone question), return it EXACTLY as-is, unchanged.

    New Question:
    {question}
    History:
    {history_text}
    """
    response = requests.post(
        "http://localhost:11434/api/generate",
        json = {
            "model": "qwen3:1.7b",
            "prompt": prompt,
            "stream": False,
            "think": False,
            "options": {"temperature": 0.2}
        }
    )

    return response.json()["response"].strip()

@app.get("/")
async def root():
    return FileResponse("static/index.html")

app.mount("/static", StaticFiles(directory="static"), name="static")

class Question(BaseModel):
    question: str
    session_id: str | None = None
    n_results: int = 3
    temperature: float = 0.3
    top_k: int = 40

def stream_ollama_response(prompt, temperature = 0.3, top_k = 40):
    response = requests.post(
        "http://localhost:11434/api/generate",
        json={
            "model": "qwen3:1.7b",
            "prompt": prompt,
            "stream": True,
            "think": False,
            "options" : {
                "temperature": temperature,
                "top_k": top_k
            }
        },
        stream=True
    )

    for line in response.iter_lines():
        if line:
            data = json.loads(line)

            token = data.get("response", "")

            if token:
                yield token

def build_prompt(question, documents, history):
    context_block = "\n\n".join(
        f"[{i}] {doc}" for i, doc in enumerate(documents)
    )
    return f"""Answer the question using ONLY the context below.

RULES:
- Give a complete, well-explained answer using all relevant details from the context — do not just give a one-line answer if the context has more supporting information available.
- State the answer once. Do not repeat it in a different phrasing afterward, and do not write "Therefore, the answer is" or similar summary phrases.
- Do not wrap your answer in quotation marks.
- Write in third person / neutral tone, not first person.
- Every sentence that uses information from the context MUST end with a citation like [0] or [1].
- Use the number that matches the source block below.
- Example: "Python was created by Guido van Rossum [0]. It supports multiple paradigms, including procedural, object-oriented, and functional styles [1]."
- If the context doesn't contain the answer, say so once — do not make up information.

Context:
{context_block}

Question: {question}

History: {history}

Answer (remember to cite every fact with [n]):"""

def generate_stream(prompt, history, question, session_id, citations, temperature, top_k):
    full_answer = ""

    for token in stream_ollama_response(prompt, temperature, top_k):
        full_answer += token
        yield token

    history.append((question, full_answer))
    chat_history[session_id] = history

    matches = re.findall(r"\[(\d+)\]", full_answer)
    used_indices = set()
    for match in matches:
        used_indices.add(int(match))

    final_citations = []
    for i, citation in enumerate(citations):
        if i in used_indices:
            final_citations.append(citation)

    print("used_indices:", used_indices)
    print("final_citations:", final_citations)

    yield "\n\n__CITATIONS_START__\n"
    yield json.dumps(final_citations)

@app.post("/chat")
def chat(q: Question):
    session_id = q.session_id or "default"
    history = chat_history.get(session_id, [])
    rewritten_question = rewrite_query(q.question, history)
    question_vector = get_embedding_single(rewritten_question)

    results = collection.query(query_embeddings=[question_vector], n_results=q.n_results)

    THRESHOLD = 0.65
    scored = []
    
    for document, distance, metadata in zip(results["documents"][0], results["distances"][0], results["metadatas"][0]):
        if distance < THRESHOLD:
            scored.append((document, distance, metadata))

    scored.sort(key=lambda x: x[1])

    for document, distance, metadata in scored:
            print(f"Page {metadata.get('page')}: {document[:100]}")

    filtered_documents = []
    for document, distance, metadata in scored:
        filtered_documents.append(document)

    if not filtered_documents:
        return {"answer": "No relevant information found in the uploaded document."}

    citations = []
    for document, distance, metadata in scored:
        citation = ({
            "source": metadata["filename"],
            "chunk_index": metadata["chunk_index"],
            "distance": round(distance, 4),
            "context": document[:250] + ("..." if len(document) > 250 else "")
        })
        if "page" in metadata:
            citation["page"] = metadata["page"]
        if "heading" in metadata:
            citation["heading"] = metadata["heading"]

        citations.append(citation)
    for i, c in enumerate(citations):
        print(f"citations[{i}] -> page {c.get('page')}, chunk_index {c.get('chunk_index')}")
    prompt = build_prompt(q.question, filtered_documents, history)

    return StreamingResponse(
        generate_stream(prompt, history, q.question, session_id, citations, q.temperature, q.top_k),
        media_type="text/plain"
    )

def ResetRequest(BaseModel):
    session_id: str

@app.post("/chat/reset")
def reset_chat(req: ResetRequest):
    chat_history.pop(req.session_id, None)
    return {"status": "reset", "Session_ID": req.session_id}

def is_pdf_heading_line(line_text, avg_size, fontnames, body_size):
    if not line_text.strip():
        return False
    if len(line_text) > 100:
        return False
    if line_text.strip().endswith("."):
        return False
    is_bigger = avg_size > body_size + 1.5
    is_bold_font = False
    for fname in fontnames:
        if "Bold" in fname:
            is_bold_font = True
    return is_bigger or is_bold_font

#to extract text from a pdf file
def pdf_extract(file_bytes):
    pages = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for i, page in enumerate(pdf.pages):
            page_text = page.extract_text() 
            if page_text:
                pages.append((i + 1, page_text))
    return pages

def pdf_extract_with_headings(file_bytes):
    all_sizes = []
    pages_data = []

    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for i, page in enumerate(pdf.pages):
            words = page.extract_words(extra_attrs=["size", "fontname"])
            pages_data.append((i + 1, words))
            for w in words:
                all_sizes.append(round(w["size"]))

    if not all_sizes:
        return []

    size_counts = Counter(all_sizes)
    body_size = size_counts.most_common(1)[0][0]

    current_heading = ""
    segments = []

    for page_num, words in pages_data:
        lines = {}
        for w in words:
            key = round(w["top"])
            if key not in lines:
                lines[key] = []
            lines[key].append(w)

        sorted_tops = sorted(lines.keys())

        for top in sorted_tops:
            line_words = lines[top]
            line_words_sorted = sorted(line_words, key=lambda w: w["x0"])

            text_parts = []
            for w in line_words_sorted:
                text_parts.append(w["text"])
            line_text = " ".join(text_parts)

            total_size = 0
            for w in line_words_sorted:
                total_size += w["size"]
            avg_size = total_size / len(line_words_sorted)

            fontnames = []
            for w in line_words_sorted:
                fontnames.append(w.get("fontname", ""))

            if is_pdf_heading_line(line_text, avg_size, fontnames, body_size):
                current_heading = line_text.strip()
                continue

            if line_text.strip():
                segments.append((line_text.strip(), page_num, current_heading))

    return segments

def chunking_pdf_with_headings(segments, chunk_size=500, overlap_sentence=1):
    grouped = []
    for text, page_num, heading in segments:
        key = (page_num, heading)
        if grouped and grouped[-1][0] == key:
            grouped[-1][1].append(text)
        else:
            grouped.append((key, [text]))

    result = []
    for key, texts in grouped:
        page_num, heading = key
        combined = " ".join(texts)
        for c in chunking(combined, chunk_size, overlap_sentence):
            result.append((c, page_num, heading))
    return result

def looks_like_heading(paragraph):
    text = paragraph.text.strip()
    if not text or len(text) > 80:
        return False
    if text.endswith("."):
        return False
    if not paragraph.runs:
        return False
    first_run = paragraph.runs[0]
    if first_run.bold:
        return True
    return False

#to extract text from a docx file
def extract_text_from_docx(file_bytes):
    stream = io.BytesIO(file_bytes)
    doc = Document(stream)
    text = ""
    #plain text extraction with heading preservation
    heading_levels = {"Heading 1": 1, "Heading 2": 2, "Heading 3": 3, "Heading 4": 4}
    current = {1: "", 2: "", 3: "", 4: ""}
    segments = []
    current_heading_flat = ""

    for paragraph in doc.paragraphs:
        para_text = paragraph.text.strip()
        if not para_text:
            continue
        style_name = paragraph.style.name

        if style_name in heading_levels:
            level = heading_levels[style_name]
            current[level] = para_text
            for i in range(level + 1, 5):
                current[i] = ""
            current_heading_flat = para_text
            continue

        if looks_like_heading(paragraph):
            current_heading_flat = para_text
            for i in range(1, 5):
                current[i] = ""
            continue

        heading_path = " > ".join(h for h in current.values() if h)
        if not heading_path:
            heading_path = current_heading_flat
        segments.append((para_text, heading_path))
        
    #text extraction from table
    for table in doc.tables:
            for row in table.rows:
                fin = []
                for cell in row.cells:
                    fin.append(cell.text)
                fin_res = "|".join(fin)
                if fin_res.strip():
                    segments.append((fin_res, "Table"))
    
    return segments

@app.get("/documents")
def list_documents():
    metadatas = collection.get()["metadatas"]
    filenames = list({m["filename"] for m in metadatas if m})

    return {"documents": filenames}

def clean_whitespace(text):
    text = re.sub(r"[ \t]+", " ", text)      # collapse multiple spaces/tabs into one
    text = re.sub(r"\n{3,}", "\n\n", text)   # collapse 3+ newlines into just 2
    return text.strip()

def decode_text_file(content):
            encodings_to_try = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
            for encoding in encodings_to_try:
                try:
                    return content.decode(encoding)
                except UnicodeDecodeError:
                    continue
            return content.decode("utf-8", errors="replace")

def chunking_with_headings(segments, chunk_size=500, overlap_sentence=1):
    grouped = []
    for text, heading in segments:
        if grouped and grouped[-1][0] == heading:
            grouped[-1][1].append(text)
        else:
            grouped.append((heading, [text]))

    result = []
    for heading, texts in grouped:
        combined = " ".join(texts)
        for c in chunking(combined, chunk_size, overlap_sentence):
            result.append((c, heading))
    return result

@app.post("/upload")
async def upload(file: UploadFile):
    MAX_FILE_SIZE = 10 * 1024 * 1024
    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="File too large. Maximum allowed size is 10MB.")
    page_chunks = []
    if file.filename.endswith(".pdf"):
        segments = pdf_extract_with_headings(content)
        cleaned = []
        for t, p, h in segments:
            cleaned.append((clean_whitespace(t), p, h))
        for c, p, h in chunking_pdf_with_headings(cleaned):
            page_chunks.append((c, p, h))
        
    elif file.filename.endswith(".txt"):
        text = decode_text_file(content)
        text = clean_whitespace(text)
        chunks = chunking(text)
        for c in chunks:
            page_chunks.append((c,None, None))
        
    elif file.filename.endswith(".docx"):
        segments = extract_text_from_docx(content)
        cleaned = []
        for t, h in segments:
            cleaned.append((clean_whitespace(t), h))
        for c, h in chunking_with_headings(cleaned):
            page_chunks.append((c, None, h))

    else:
        raise HTTPException(status_code=400, detail="Unsupported file type. Please upload a PDF, TXT, or DOCX file.")
    
    collection.delete(where={"filename": file.filename})

    filtered = []
    for c, p, h in page_chunks:
        if c.strip() and len(c.strip()) > 3:
            filtered.append((c, p, h))
    page_chunks = filtered

    if not page_chunks:
        raise HTTPException(status_code=400, detail="No usable content extracted from this file.")

    chunk_texts = []
    for c, p, h in page_chunks:
        chunk_texts.append(c)

    start = time.time()#remove later
    vectors = get_embeddings_batch(chunk_texts)  
    print(f"Embedded {len(chunk_texts)} chunks in {time.time() - start:.2f}s")#remove later

    metadatas = []
    for i, (chunk_text, page_num, heading) in enumerate(page_chunks):
        meta = {"filename": file.filename, "chunk_index": i}
        if page_num is not None:
            meta["page"] = page_num
        if heading:
            meta["heading"] = heading
        metadatas.append(meta)
    
    collection.add(
        ids=[str(uuid.uuid4()) for _ in page_chunks],
        embeddings=vectors,
        documents=chunk_texts,
        metadatas=metadatas
    )

    return {"filename": file.filename, "chunks_stored": len(page_chunks)}

class DeleteRequest(BaseModel):
    filename:str

@app.post("/delete")
def delete(q:DeleteRequest):
    collection.delete(
        where = {"filename": q.filename}
    )
    return {"Status":"Deleted", "filename": q.filename}

class URLingest(BaseModel):
    url: str

@app.post("/ingest-url")
def ingest_url(req: URLingest):
    validate_url(req.url)
    session = requests.Session()
    session.headers.update({"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"})
    
    try:
        response = session.get(req.url, timeout = 10)
    
        response.raise_for_status()
    except requests.exceptions.RequestException as e:
        raise HTTPException(status_code=400, detail=f"Failed to fetch URL: {e}")
    
    page_chunks = []
    if req.url.endswith(".pdf") or "application/pdf" in response.headers.get("Content-Type", ""):
        pages = pdf_extract(response.content)          
        for page_num, page_text in pages:
            page_text = clean_whitespace(page_text)
            chunks = chunking(page_text)
            for c in chunks:
                page_chunks.append((c, page_num))
    else:
        soup = BeautifulSoup(response.text, "lxml")
        text = soup.get_text(separator="\n")
        text = clean_whitespace(text)
        chunks = chunking(text)
        for c in chunks:
            page_chunks.append((c, None))

    page_chunks = [(c, p) for c, p in page_chunks if c.strip() and len(c.strip()) > 3]

    if not page_chunks:
        raise HTTPException(400, "No usable content after chunking.")

    filename = req.url
    collection.delete(where={"filename": filename})

    chunk_texts = [c for c, p in page_chunks]

    vectors = get_embeddings_batch(chunk_texts)
    metadatas = []
    for i, (chunk_text, page_num) in enumerate(page_chunks):
        meta = {"filename": filename, "chunk_index": i}
        if page_num is not None:
            meta["page"] = page_num
        metadatas.append(meta)

    collection.add(
        ids=[str(uuid.uuid4()) for _ in chunk_texts],
        embeddings=vectors,
        documents=chunk_texts,
        metadatas=metadatas
    )

    return {"filename": filename, "chunks_stored": len(chunk_texts)}

PRIVATE_NETS = [
    ipaddress.ip_network("10.0.0.0/8"),
    ipaddress.ip_network("172.16.0.0/12"),
    ipaddress.ip_network("192.168.0.0/16"),
    ipaddress.ip_network("127.0.0.0/8"),
    ipaddress.ip_network("169.254.0.0/16"),
]

def is_private(hostname):
    try:
        ip = ipaddress.ip_address(hostname)
    except ValueError:
        ip = ipaddress.ip_address(socket.gethostbyname(hostname))
    return any(ip in net for net in PRIVATE_NETS)

def validate_url(url):
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https"):
        raise HTTPException(400, "Only http/https URLs allowed")
    if not parsed.hostname or is_private(parsed.hostname):
        raise HTTPException(400, "The given ip is not allowed")